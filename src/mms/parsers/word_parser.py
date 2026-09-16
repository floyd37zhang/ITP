from __future__ import annotations

from typing import Optional

from docx import Document

from lib.logger import get_logger

from ..models.question import Question
from .base import FileParser, extract_questions_from_text

log = get_logger("mms.parsers.word")


class WordParser(FileParser):

    def parse(self, file_path: str) -> str:
        try:
            doc = Document(file_path)
            parts: list[str] = []

            for para in doc.paragraphs:
                if para.text.strip():
                    parts.append(para.text)

            for table in doc.tables:
                for row in table.rows:
                    cells = [cell.text.strip() for cell in row.cells]
                    line = " | ".join(c for c in cells if c)
                    if line:
                        parts.append(line)

            return "\n".join(parts)
        except Exception as exc:
            log.error("Failed to parse Word file %s: %s", file_path, exc)
            return ""

    def extract_questions(self, file_path: str) -> list[Question]:
        text = self.parse(file_path)
        if not text:
            return []
        return extract_questions_from_text(text)
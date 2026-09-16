import sys, os, tempfile
sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

from fastapi.testclient import TestClient
from docx import Document

from main import _build_app
from config import AppConfig
from storage.memory_storage import MemoryStorage
from storage.mysql_storage import MySQLStorage
from models.resource import Resource, ResourceQuery, ResourceType
from models.question import Question, QuestionType
from parsers.base import extract_questions_from_text
from parsers.word_parser import WordParser
from scanner.file_scanner import (
    FileScanner, compute_file_hash, infer_metadata_from_path, guess_resource_type,
)

passed = 0
failed = 0


def check(name: str, cond: bool) -> None:
    global passed, failed
    if cond:
        print(f"  PASS  {name}")
        passed += 1
    else:
        print(f"  FAIL  {name}")
        failed += 1


def test_models():
    print("\n[TEST] Models")
    r = Resource(
        file_path="/data/六年级/数学/人教版/期末试卷.docx",
        file_name="期末试卷.docx",
        file_extension=".docx",
        grade="六年级", subject="数学", version="人教版",
        resource_type=ResourceType.EXAM_PAPER,
        knowledge_points=["分数运算"], difficulty=3,
        class_level="普通班", focus_tags=["计算能力"],
        file_size=10240, file_hash="abc123",
    )
    d = r.to_dict()
    check("resource.to_dict grade", d["grade"] == "六年级")
    check("resource.to_dict type", d["resource_type"] == "exam_paper")
    check("resource.to_dict list", d["knowledge_points"] == ["分数运算"])

    q = Question(
        resource_id=1, order_index=1,
        question_type=QuestionType.SINGLE_CHOICE,
        content="1+1=?", options=["A.1", "B.2", "C.3"],
        answer="B", score=5.0,
    )
    check("question.type", q.question_type == QuestionType.SINGLE_CHOICE)
    check("question.score", q.score == 5.0)


def test_memory_storage():
    print("\n[TEST] MemoryStorage")
    storage = MemoryStorage()

    r = Resource(
        file_path="/a/b/c.docx", file_name="c.docx", file_extension=".docx",
        grade="五年级", subject="语文", version="部编版",
        resource_type=ResourceType.LESSON_PLAN,
        knowledge_points=["古诗词", "文言文"], difficulty=2,
        class_level="重点班", focus_tags=["朗读训练"],
        file_size=5000, file_hash="h1",
    )
    saved = storage.save_resource(r)
    check("save returns id", saved.id is not None)
    rid = saved.id

    fetched = storage.get_resource(rid)
    check("get by id", fetched is not None and fetched.grade == "五年级")

    by_path = storage.get_resource_by_path("/a/b/c.docx")
    check("get by path", by_path is not None and by_path.id == rid)

    items, total = storage.query_resources(ResourceQuery(grade="五年级", limit=10))
    check("query by grade total", total == 1)

    items2, _ = storage.query_resources(ResourceQuery(subject="数学"))
    check("query empty", len(items2) == 0)

    items3, _ = storage.query_resources(ResourceQuery(keyword="古诗"))
    check("query keyword", len(items3) == 1)

    ok = storage.delete_resource(rid)
    check("delete", ok is True)
    check("after delete gone", storage.get_resource(rid) is None)


def test_parsers():
    print("\n[TEST] Parsers")
    sample = """
一、单项选择题
1. 下列哪个数是质数？
A. 4
B. 6
C. 7
D. 9

二、填空题
2. 2 + 3 = ______

三、计算题
3. 解方程：2x + 5 = 15
"""
    questions = extract_questions_from_text(sample)
    check("extract count", len(questions) >= 2)
    check("q1 has content", questions[0].content != "")

    types = [q.question_type for q in questions]
    check("has single choice", QuestionType.SINGLE_CHOICE in types)


def test_scanner():
    print("\n[TEST] FileScanner")

    with tempfile.TemporaryDirectory() as tmpdir:
        path = os.path.join(tmpdir, "六年级", "数学", "人教版")
        os.makedirs(path)

        doc = Document()
        doc.add_paragraph("一、选择题")
        doc.add_paragraph("1. 计算 1/2 + 1/3 = ?")
        doc.add_paragraph("A. 5/6")
        doc.add_paragraph("B. 1/5")
        doc.add_paragraph("C. 2/5")
        doc.add_paragraph("D. 3/5")
        doc.save(os.path.join(path, "期末试卷.docx"))

        storage = MemoryStorage()
        scanner = FileScanner(
            storage=storage,
            parsers={".docx": WordParser()},
            supported_extensions=(".docx",),
            scan_paths=[tmpdir],
            scan_interval=3600,
        )

        result = scanner.scan_once()
        check("scanner scanned >= 1", result["scanned"] >= 1)

        resources, _ = storage.query_resources(ResourceQuery(limit=100))
        check("resource indexed", len(resources) >= 1)
        r = resources[0]
        check("grade inferred", r.grade == "六年级")
        check("subject inferred", r.subject == "数学")
        check("type = exam", r.resource_type == ResourceType.EXAM_PAPER)

        questions = storage.get_questions_by_resource(r.id or 0)
        check("questions extracted", len(questions) >= 1)

        # Test file deletion detection
        os.remove(os.path.join(path, "期末试卷.docx"))
        scanner.scan_once()
        resources2, _ = storage.query_resources(ResourceQuery(limit=100))
        check("file deleted -> resource removed", len(resources2) == 0)


def test_metadata_inference():
    print("\n[TEST] Metadata Inference")
    meta = infer_metadata_from_path(
        "/data/五年级/数学/人教版/重点班/教案1.docx"
    )
    check("grade", meta.get("grade") == "五年级")
    check("subject", meta.get("subject") == "数学")
    check("version", meta.get("version") == "人教版")
    check("class_level", meta.get("class_level") == "重点班")

    check("exam paper", guess_resource_type("期末试卷.docx", ".docx") == ResourceType.EXAM_PAPER)
    check("lesson plan", guess_resource_type("教案.docx", ".docx") == ResourceType.LESSON_PLAN)
    check("image", guess_resource_type("photo.jpg", ".jpg") == ResourceType.IMAGE)


def test_api():
    print("\n[TEST] FastAPI API")
    cfg = AppConfig()
    cfg.storage_backend = "memory"
    app, scanner, _ = _build_app(cfg)
    client = TestClient(app)

    r = client.get("/api/v1/admin/health")
    check("health status", r.status_code == 200)
    check("health ok", r.json()["status"] == "ok")

    payload = {
        "file_path": "/test/六年级/数学/人教版/教案1.docx",
        "grade": "六年级", "subject": "数学", "version": "人教版",
        "resource_type": "lesson_plan",
        "knowledge_points": ["分数运算"],
        "difficulty": 3, "class_level": "重点班",
        "focus_tags": ["概念讲解"],
    }
    r = client.post("/api/v1/resources", json=payload)
    check("create resource", r.status_code == 200)
    created = r.json()
    rid = created["id"]

    r = client.get(f"/api/v1/resources/{rid}")
    check("get resource", r.status_code == 200)
    check("grade matches", r.json()["grade"] == "六年级")

    r = client.get("/api/v1/resources?subject=数学")
    check("search by subject", r.status_code == 200)
    check("total >= 1", r.json()["total"] >= 1)

    r = client.post("/api/v1/questions", json=[{
        "resource_id": rid, "order_index": 1,
        "question_type": "single_choice",
        "content": "1+1=?", "options": ["1", "2", "3", "4"],
        "answer": "B", "score": 5.0,
    }])
    check("create question", r.status_code == 200)

    r = client.get(f"/api/v1/questions/by-resource/{rid}")
    check("get questions", r.status_code == 200)
    check("question count", len(r.json()["questions"]) >= 1)

    r = client.delete(f"/api/v1/resources/{rid}")
    check("delete", r.status_code == 200)
    check("success", r.json()["success"] is True)

    r = client.get("/api/v1/resources/99999")
    check("404 not found", r.status_code == 404)


def main():
    test_models()
    test_memory_storage()
    test_metadata_inference()
    test_parsers()
    test_scanner()
    test_api()

    print(f"\n{'='*40}")
    print(f"  PASSED: {passed}  |  FAILED: {failed}")
    print(f"{'='*40}")
    return failed == 0


if __name__ == "__main__":
    ok = main()
    sys.exit(0 if ok else 1)